"""Bounded extraction and lexical ranking for uploaded study material."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import logging
import re
from zipfile import BadZipFile, ZipFile

import fitz
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.config import settings
from app.services.embedding_service import cosine_similarity, embed_texts

logger = logging.getLogger(__name__)

MAX_SOURCE_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 5_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PAGE_COUNT = 500
MAX_EXTRACTED_CHARACTERS = 2_000_000
MAX_CONTENT_UNITS = 10_000
MAX_STUDY_FIGURES = 12
MAX_FIGURE_PIXELS = 32_000_000
CHUNK_CHARACTERS = 1_200
CHUNK_OVERLAP = 200
CHILD_CHUNK_CHARACTERS = 850
PARENT_CHUNK_CHARACTERS = 2_400
MIN_SEMANTIC_PARENT_CHARACTERS = 350
SEMANTIC_BOUNDARY_THRESHOLD = 0.55
CHILD_SEMANTIC_BOUNDARY_THRESHOLD = 0.42

PDF_CONTENT_TYPE = "application/pdf"
PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_TOKEN_PATTERN = re.compile(r"[\w]+", re.UNICODE)
_SENTENCE_BOUNDARY = re.compile(r"(?<=[.!?])\s+")


class StudyMaterialError(ValueError):
    """Raised when a study file is unsupported, invalid, or exceeds safe limits."""


class StudyEmbeddingUnavailable(RuntimeError):
    """Raised when semantic chunking cannot reach its configured local embedding model."""


@dataclass(frozen=True, slots=True)
class ContentUnit:
    location_type: str
    location_number: int
    chunk_index: int
    content: str
    chunk_level: str = "child"
    parent_chunk_index: int | None = None
    char_start: int | None = None
    char_end: int | None = None


def tokenize(text: str) -> list[str]:
    return [token.casefold() for token in _TOKEN_PATTERN.findall(text) if len(token) > 1]


def chunk_text(text: str) -> list[str]:
    normalized = " ".join(text.split())
    if not normalized:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + CHUNK_CHARACTERS, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", start + 1, end)
            if boundary - start >= CHUNK_CHARACTERS // 2:
                end = boundary
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def _split_oversized_sentence(sentence: str) -> list[str]:
    pieces: list[str] = []
    remaining = sentence.strip()
    while len(remaining) > CHILD_CHUNK_CHARACTERS:
        boundary = remaining.rfind(" ", 0, CHILD_CHUNK_CHARACTERS + 1)
        if boundary < CHILD_CHUNK_CHARACTERS // 2:
            boundary = CHILD_CHUNK_CHARACTERS
        pieces.append(remaining[:boundary].strip())
        remaining = remaining[boundary:].strip()
    if remaining:
        pieces.append(remaining)
    return pieces


def _sentences(text: str) -> list[str]:
    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n+", text) if paragraph.strip()]
    sentences: list[str] = []
    for paragraph in paragraphs:
        normalized = " ".join(paragraph.split())
        for sentence in _SENTENCE_BOUNDARY.split(normalized):
            sentences.extend(_split_oversized_sentence(sentence))
    return sentences


def semantic_parent_child_chunks(
    text: str,
    location_type: str,
    location_number: int,
    *,
    embedder=None,
    child_index_start: int = 0,
    parent_index_start: int = 0,
) -> list[ContentUnit]:
    """Split by semantic boundaries, retaining contextual parents for searchable children."""
    sentences = _sentences(text)
    if not sentences:
        return []

    encode = embedder or embed_texts
    try:
        vectors = encode(sentences)
    except Exception as exc:
        raise StudyEmbeddingUnavailable(
            "The semantic embedding model could not process this study source"
        ) from exc
    if len(vectors) != len(sentences):
        raise RuntimeError("Semantic embedding model returned an unexpected number of sentences")

    parents: list[list[int]] = []
    current: list[int] = []
    current_characters = 0
    for index, sentence in enumerate(sentences):
        sentence_size = len(sentence)
        semantic_break = (
            bool(current)
            and current_characters >= MIN_SEMANTIC_PARENT_CHARACTERS
            and cosine_similarity(vectors[current[-1]], vectors[index])
            < SEMANTIC_BOUNDARY_THRESHOLD
        )
        size_break = bool(current) and current_characters + sentence_size + 1 > PARENT_CHUNK_CHARACTERS
        if semantic_break or size_break:
            parents.append(current)
            current = []
            current_characters = 0
        current.append(index)
        current_characters += sentence_size + 1
    if current:
        parents.append(current)

    units: list[ContentUnit] = []
    next_child_index = child_index_start
    next_parent_index = parent_index_start
    for parent_sentences in parents:
        parent_index = next_parent_index
        next_parent_index += 1
        parent_text = " ".join(sentences[index] for index in parent_sentences)
        units.append(
            ContentUnit(
                location_type,
                location_number,
                parent_index,
                parent_text,
                chunk_level="parent",
            )
        )

        child_groups: list[list[int]] = []
        child_group: list[int] = []
        child_characters = 0
        for sentence_index in parent_sentences:
            sentence_size = len(sentences[sentence_index])
            semantic_break = (
                bool(child_group)
                and child_characters >= CHILD_CHUNK_CHARACTERS // 2
                and cosine_similarity(
                    vectors[child_group[-1]],
                    vectors[sentence_index],
                )
                < CHILD_SEMANTIC_BOUNDARY_THRESHOLD
            )
            size_break = (
                bool(child_group)
                and child_characters + sentence_size + 1 > CHILD_CHUNK_CHARACTERS
            )
            if semantic_break or size_break:
                child_groups.append(child_group)
                child_group = []
                child_characters = 0
            child_group.append(sentence_index)
            child_characters += sentence_size + 1
        if child_group:
            child_groups.append(child_group)

        parent_offset = 0
        for child_group in child_groups:
            child_text = " ".join(sentences[index] for index in child_group)
            char_start = parent_offset
            char_end = char_start + len(child_text)
            units.append(
                ContentUnit(
                    location_type,
                    location_number,
                    next_child_index,
                    child_text,
                    chunk_level="child",
                    parent_chunk_index=parent_index,
                    char_start=char_start,
                    char_end=char_end,
                )
            )
            next_child_index += 1
            parent_offset = char_end + 1

    return units


def _validate_pptx_archive(file_bytes: bytes) -> None:
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise StudyMaterialError("PowerPoint archive contains too many files")
            if any(info.flag_bits & 0x1 for info in members):
                raise StudyMaterialError("Encrypted PowerPoint files are not supported")
            if any(
                name.startswith(("/", "\\")) or ".." in name.replace("\\", "/").split("/")
                for name in (info.filename for info in members)
            ):
                raise StudyMaterialError("PowerPoint archive contains an invalid path")
            if sum(info.file_size for info in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise StudyMaterialError("PowerPoint archive expands beyond the allowed size")
            if any(
                info.file_size > 0
                and (info.compress_size == 0 or info.file_size / info.compress_size > 200)
                for info in members
            ):
                raise StudyMaterialError("PowerPoint archive has an unsafe compression ratio")
            names = {info.filename for info in members}
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(names):
                raise StudyMaterialError("The uploaded file is not a valid PowerPoint presentation")
    except BadZipFile as exc:
        raise StudyMaterialError("The uploaded file is not a valid PowerPoint presentation") from exc


def _pdf_pages(file_bytes: bytes) -> list[tuple[int, str, list[bytes]]]:
    if b"%PDF-" not in file_bytes[:1024]:
        raise StudyMaterialError("The uploaded file does not have a valid PDF signature")
    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
    except (fitz.FileDataError, RuntimeError, ValueError) as exc:
        raise StudyMaterialError("The uploaded PDF is invalid or unreadable") from exc

    try:
        if document.is_encrypted:
            raise StudyMaterialError("Password-protected PDFs are not supported")
        if document.page_count > MAX_PAGE_COUNT:
            raise StudyMaterialError(f"PDF exceeds the {MAX_PAGE_COUNT}-page limit")
        pages = []
        for page_number in range(1, document.page_count + 1):
            page = document.load_page(page_number - 1)
            images = []
            seen_xrefs = set()
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                extracted = document.extract_image(xref)
                image_bytes = extracted.get("image", b"")
                if image_bytes:
                    images.append(image_bytes)
            pages.append((page_number, page.get_text("text"), images))
        return pages
    except StudyMaterialError:
        raise
    except (fitz.FileDataError, RuntimeError, ValueError) as exc:
        raise StudyMaterialError("The uploaded PDF is invalid or unreadable") from exc
    finally:
        document.close()


def _pptx_slides(file_bytes: bytes) -> list[tuple[int, str, list[bytes]]]:
    _validate_pptx_archive(file_bytes)
    try:
        presentation = Presentation(BytesIO(file_bytes))
        if len(presentation.slides) > MAX_PAGE_COUNT:
            raise StudyMaterialError(f"Presentation exceeds the {MAX_PAGE_COUNT}-slide limit")
        slides: list[tuple[int, str, list[bytes]]] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            paragraphs: list[str] = []
            images: list[bytes] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    paragraphs.append(shape.text)
                if shape.has_table:
                    paragraphs.extend(
                        cell.text
                        for row in shape.table.rows
                        for cell in row.cells
                    )
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append(shape.image.blob)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    paragraphs.append(f"[Speaker notes]\n{notes}")
            slides.append((slide_number, "\n".join(paragraphs), images))
        return slides
    except StudyMaterialError:
        raise
    except Exception as exc:
        raise StudyMaterialError("The uploaded PowerPoint file is invalid or unreadable") from exc


def _prepare_figure_image(image_bytes: bytes) -> bytes:
    try:
        with Image.open(BytesIO(image_bytes)) as image:
            if image.width * image.height > MAX_FIGURE_PIXELS:
                raise StudyMaterialError("A study figure exceeds the 32-megapixel processing limit")
            image.thumbnail((1_400, 1_400))
            normalized = image.convert("RGB")
            output = BytesIO()
            normalized.save(output, format="JPEG", quality=84, optimize=True)
            return output.getvalue()
    except StudyMaterialError:
        raise
    except (OSError, UnidentifiedImageError) as exc:
        raise StudyMaterialError("A study figure image is invalid or unreadable") from exc


def _describe_figures(images: list[bytes], location_type: str, location_number: int) -> list[str]:
    if not images:
        return []
    vision_configured = bool(settings.GEMINI_API_KEY.get_secret_value())
    if vision_configured:
        from app.services.vlm_service import describe_study_figure
    from app.services.ocr_services import extract_text_from_image

    descriptions = []
    for index, image_bytes in enumerate(images, start=1):
        try:
            image = _prepare_figure_image(image_bytes)
            if vision_configured:
                try:
                    description = describe_study_figure(image, "image/jpeg")
                    source_label = "VLM figure description"
                except Exception:
                    logger.exception(
                        "VLM failed; attempting OCR for %s %s figure %s",
                        location_type,
                        location_number,
                        index,
                    )
                    ocr_text = extract_text_from_image(image)
                    if not ocr_text:
                        raise
                    description = ocr_text
                    source_label = "OCR fallback (VLM unavailable)"
            else:
                description = extract_text_from_image(image)
                source_label = "OCR figure text"
            if not description.strip():
                logger.warning(
                    "Study figure at %s %s yielded no description or readable OCR text",
                    location_type,
                    location_number,
                )
                continue
        except Exception as exc:
            logger.exception(
                "Study figure analysis failed at %s %s figure %s",
                location_type,
                location_number,
                index,
            )
            continue
        descriptions.append(
            f"[Figure {index} on {location_type} {location_number}; {source_label}]\n{description}"
        )
    return descriptions


def extract_study_material(file_bytes: bytes, filename: str) -> tuple[str, int, list[ContentUnit]]:
    if not file_bytes:
        raise StudyMaterialError("The uploaded file is empty")
    if len(file_bytes) > MAX_SOURCE_BYTES:
        raise StudyMaterialError("The uploaded file exceeds the 25 MB limit")

    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if extension == "pdf":
        content_type = PDF_CONTENT_TYPE
        locations = _pdf_pages(file_bytes)
        location_type = "page"
    elif extension == "pptx":
        content_type = PPTX_CONTENT_TYPE
        locations = _pptx_slides(file_bytes)
        location_type = "slide"
    else:
        raise StudyMaterialError("Only PDF and PowerPoint (.pptx) files are supported")

    units: list[ContentUnit] = []
    extracted_characters = 0
    figure_count = 0
    for location_number, text, images in locations:
        figure_count += len(images)
        if figure_count > MAX_STUDY_FIGURES:
            raise StudyMaterialError(
                f"The document contains more than the {MAX_STUDY_FIGURES}-figure processing limit"
            )
        figure_descriptions = _describe_figures(images, location_type, location_number)
        combined_text = "\n\n".join([text, *figure_descriptions])
        extracted_characters += len(combined_text)
        if extracted_characters > MAX_EXTRACTED_CHARACTERS:
            raise StudyMaterialError("Extracted text exceeds the 2 MB processing limit")
        location_units = semantic_parent_child_chunks(
            combined_text,
            location_type,
            location_number,
        )
        units.extend(location_units)
        if len(units) > MAX_CONTENT_UNITS:
            raise StudyMaterialError("The document creates too many parent and child chunks")

    if not units:
        raise StudyMaterialError(
            "No searchable text or configured VLM-readable figures were found in the document."
        )
    return content_type, len(locations), units


def rank_content(query: str, content: str) -> float:
    query_terms = set(tokenize(query))
    if not query_terms:
        return 0.0

    term_counts = Counter(tokenize(content))
    matched_terms = query_terms.intersection(term_counts)
    if not matched_terms:
        return 0.0

    score = sum(min(term_counts[term], 3) for term in matched_terms) / len(query_terms)
    if query.casefold().strip() in content.casefold():
        score += 1.0
    return score
