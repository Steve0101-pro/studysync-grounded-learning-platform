from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from app.services import study_materials
from app.services.study_materials import (
    CHUNK_CHARACTERS,
    MAX_SOURCE_BYTES,
    StudyMaterialError,
    chunk_text,
    extract_study_material,
    rank_content,
    semantic_parent_child_chunks,
)


def _make_pdf(*page_texts: str) -> bytes:
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), text)
    result = document.tobytes()
    document.close()
    return result


def _make_pptx(slide_texts: list[str], slide_notes: list[str] | None = None) -> bytes:
    presentation = Presentation()
    for index, text in enumerate(slide_texts):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text_frame.text = text
        if slide_notes:
            slide.notes_slide.notes_text_frame.text = slide_notes[index]
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.fixture(autouse=True)
def deterministic_semantic_embeddings(monkeypatch):
    monkeypatch.setattr(
        study_materials,
        "embed_texts",
        lambda texts: [[1.0, 0.0] for _ in texts],
    )


def test_pdf_extraction_keeps_page_citations():
    content_type, page_count, units = extract_study_material(
        _make_pdf("Photosynthesis converts light into chemical energy.", "Chlorophyll absorbs light."),
        "biology.pdf",
    )

    assert content_type == "application/pdf"
    assert page_count == 2
    children = [unit for unit in units if unit.chunk_level == "child"]
    assert [(unit.location_type, unit.location_number) for unit in children] == [
        ("page", 1),
        ("page", 2),
    ]
    assert "Photosynthesis" in children[0].content
    assert all(unit.chunk_level in {"parent", "child"} for unit in units)
    assert all(unit.parent_chunk_index is not None for unit in children)


def test_pptx_extraction_keeps_slide_citations():
    content_type, page_count, units = extract_study_material(
        _make_pptx(["Cell structure", "Mitochondria produce ATP."]),
        "biology.pptx",
    )

    assert content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert page_count == 2
    children = [unit for unit in units if unit.chunk_level == "child"]
    assert [(unit.location_type, unit.location_number) for unit in children] == [
        ("slide", 1),
        ("slide", 2),
    ]
    assert "Mitochondria" in children[1].content


def test_pptx_speaker_notes_are_added_to_searchable_content():
    _, _, units = extract_study_material(
        _make_pptx(["Cell structure"], ["Mitochondria produce ATP during respiration."]),
        "biology.pptx",
    )

    child = next(unit for unit in units if unit.chunk_level == "child")
    assert "Speaker notes" in child.content
    assert "Mitochondria produce ATP" in child.content


def test_pdf_figures_are_described_and_indexed_with_their_page(monkeypatch):
    image_bytes = BytesIO()
    Image.new("RGB", (40, 40), color="blue").save(image_bytes, format="PNG")
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Cell biology")
    page.insert_image(fitz.Rect(72, 100, 140, 168), stream=image_bytes.getvalue())
    pdf_bytes = document.tobytes()
    document.close()
    monkeypatch.setattr(
        study_materials,
        "_describe_figures",
        lambda *_args: ["[Figure 1 visual description]\nA labeled mitochondrion diagram."],
    )

    _, page_count, units = extract_study_material(pdf_bytes, "biology.pdf")

    assert page_count == 1
    children = [unit for unit in units if unit.chunk_level == "child"]
    assert any("labeled mitochondrion diagram" in unit.content for unit in children)
    assert all(unit.location_type == "page" and unit.location_number == 1 for unit in children)


@pytest.mark.parametrize(
    ("file_bytes", "filename", "message"),
    [
        (b"", "notes.pdf", "empty"),
        (b"not a pdf", "notes.pdf", "signature"),
        (b"content", "notes.docx", "Only PDF and PowerPoint"),
    ],
)
def test_invalid_source_files_are_rejected(file_bytes, filename, message):
    with pytest.raises(StudyMaterialError, match=message):
        extract_study_material(file_bytes, filename)


def test_image_only_pdf_is_rejected_instead_of_indexed_as_empty():
    document = fitz.open()
    document.new_page()
    pdf_bytes = document.tobytes()
    document.close()

    with pytest.raises(StudyMaterialError, match="No searchable text"):
        extract_study_material(pdf_bytes, "scanned-notes.pdf")


def test_oversized_source_is_rejected_before_parsing():
    with pytest.raises(StudyMaterialError, match="25 MB limit"):
        extract_study_material(b"x" * (MAX_SOURCE_BYTES + 1), "notes.pdf")


def test_highly_compressed_powerpoint_archive_is_rejected():
    archive_bytes = BytesIO()
    with ZipFile(archive_bytes, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("ppt/presentation.xml", "x" * 100_000)

    with pytest.raises(StudyMaterialError, match="unsafe compression ratio"):
        extract_study_material(archive_bytes.getvalue(), "unsafe.pptx")


def test_text_chunks_are_bounded_and_preserve_the_source_text():
    text = " ".join(f"concept-{index}" for index in range(600))
    chunks = chunk_text(text)

    assert len(chunks) > 1
    assert all(len(chunk) <= CHUNK_CHARACTERS for chunk in chunks)
    assert all(chunk in text for chunk in chunks)


def test_semantic_chunking_persists_parent_context_and_exact_child_spans():
    text = " ".join(
        [
            "A" * 200 + ".",
            "B" * 200 + ".",
            "C" * 200 + ".",
        ]
    )
    sentence_vectors = iter([[1.0, 0.0], [1.0, 0.0], [0.0, 1.0]])
    units = semantic_parent_child_chunks(
        text,
        "page",
        1,
        embedder=lambda sentences: [next(sentence_vectors) for _ in sentences],
    )

    parents = {unit.chunk_index: unit for unit in units if unit.chunk_level == "parent"}
    children = [unit for unit in units if unit.chunk_level == "child"]

    assert len(parents) == 2
    assert len(children) == 2
    for child in children:
        parent = parents[child.parent_chunk_index]
        assert parent.content[child.char_start : child.char_end] == child.content


def test_lexical_ranking_rewards_relevant_terms_and_exact_phrases():
    exact = rank_content("mitochondria produce ATP", "Mitochondria produce ATP during respiration.")
    partial = rank_content("mitochondria produce ATP", "Mitochondria are organelles.")
    unrelated = rank_content("mitochondria produce ATP", "The cell membrane controls transport.")

    assert exact > partial > unrelated
    assert unrelated == 0
